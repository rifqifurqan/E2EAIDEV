"""Real Docling conversions (FR-D3). Slow on the first run: the layout, table and OCR models are downloaded and loaded once."""

import io
import zipfile

import pytest

pytest.importorskip("docling")

from e2eai.core.errors import AppError  # noqa: E402
from e2eai.docling_runtime import convert_bytes  # noqa: E402
from e2eai.ingest import DoclingParser, select_parser  # noqa: E402


def _pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    table = Table([["Part", "Torque (Nm)", "Notes"], ["Bolt A", "12", "Use thread locker"], ["Bolt B", "25", "Dry"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build([
        Paragraph("Partnership Agreement", styles["Title"]),
        Paragraph("1. Revenue share", styles["Heading1"]),
        Paragraph("The revenue share between the parties is 70 percent for the platform and 30 percent for the partner.", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Table 1: Torque specification", styles["Normal"]),
        table,
        PageBreak(),
        Paragraph("2. Term", styles["Heading1"]),
        Paragraph("This agreement runs for two years and renews automatically unless either party gives ninety days notice.", styles["Normal"]),
    ])
    return buf.getvalue()


def _scanned_png() -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1400, 500), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=64)
    draw.text((60, 60), "Invoice number 20481", fill="black", font=font)
    draw.text((60, 220), "Total payable 1500 dollars", fill="black", font=font)
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Employees receive twelve days of paid annual leave.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Grade", "Extra days"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Senior", "3"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture(scope="module")
def pdf_chunks():
    return convert_bytes(_pdf(), "application/pdf")


def test_a_digital_pdf_keeps_headings_page_anchors_and_text(pdf_chunks):
    revenue = next(c for c in pdf_chunks if "70 percent" in c.text)
    assert revenue.page == 1 and revenue.kind == "text" and "Revenue share" in revenue.section_path
    term = next(c for c in pdf_chunks if "two years" in c.text)
    assert term.page == 2 and "Term" in term.section_path


def test_a_pdf_table_becomes_a_structured_chunk_with_its_caption(pdf_chunks):
    table = next(c for c in pdf_chunks if c.kind == "table")
    assert table.page == 1
    assert table.table_json["rows"][0] == ["Part", "Torque (Nm)", "Notes"]
    assert ["Bolt A", "12", "Use thread locker"] in table.table_json["rows"]
    assert "Table 1" in table.table_json["caption"] and "Bolt B" in table.text


def test_every_chunk_carries_a_page_confidence_between_zero_and_one(pdf_chunks):
    scores = [c.confidence for c in pdf_chunks]
    assert all(s is not None and 0.0 <= s <= 1.0 for s in scores)


def test_a_scanned_image_is_read_with_ocr():
    text = " ".join(c.text for c in convert_bytes(_scanned_png(), "application/octet-stream")).lower()
    assert "invoice" in text and "20481" in text.replace(" ", "")  # mime is unknown: the PNG bytes identify it


def test_a_word_document_keeps_its_heading_and_table():
    chunks = convert_bytes(_docx(), "application/octet-stream")  # octet-stream: detected as .docx from the zip
    assert any("twelve days" in c.text and "Leave policy" in c.section_path for c in chunks)
    table = next(c for c in chunks if c.kind == "table")
    assert ["Senior", "3"] in table.table_json["rows"]


def test_a_damaged_pdf_is_a_clear_422_not_a_crash():
    with pytest.raises(AppError) as err:
        convert_bytes(b"%PDF-1.7 this is not really a pdf", "application/pdf")
    assert err.value.status == 422 and err.value.title == "Document could not be parsed"


def test_an_unsupported_file_is_a_clear_415():
    with pytest.raises(AppError) as err:
        convert_bytes(b"MZ\x90\x00 an executable", "application/octet-stream")
    assert err.value.status == 415


def test_a_zip_that_is_not_an_office_file_is_unsupported():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes/readme.txt", "hello")
    with pytest.raises(AppError) as err:
        convert_bytes(buf.getvalue(), "application/zip")
    assert err.value.status == 415


def test_the_ingest_parser_uses_docling_for_rich_files():
    assert isinstance(select_parser("application/pdf"), DoclingParser)
    chunks = select_parser("application/pdf").parse(_pdf(), mime="application/pdf")
    assert any("70 percent" in c.text for c in chunks)
